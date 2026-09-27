"""Build src/data/afford.json: income distribution, consumer price level and
phone/internet prices per country, for the affordability step and price anchors.

Sources (all free with attribution):
- World Bank Poverty and Inequality Platform (PIP), CC BY 4.0: share of people
  below a set of daily income lines, 2021 PPP dollars, nowcast year 2025.
- World Bank WDI, CC BY 4.0: PA.NUS.PRVT.PP (PPP factor for household
  consumption) / PA.NUS.FCRF (market exchange rate) = consumer price level vs US.
- ITU ICT Price Baskets 2008-2025, CC BY-NC-SA 3.0 IGO: mobile plan with calls
  and 5 GB, fixed broadband 5 GB, USD per month, 2025.

Run: python3 data-sources/affordability/build.py [path/to/ITU_ICTPriceBaskets_2008-2025.xlsx]
Without the xlsx argument the ITU extract in this folder is reused.
"""
import csv, io, json, pathlib, sys, time, urllib.request

HERE = pathlib.Path(__file__).parent
ROOT = HERE.parent.parent
YEAR = 2025
# Daily income lines, 2021 PPP $ per person. Dense where subscriptions bite.
LINES = [1, 2, 3, 4.5, 6, 8, 10, 13, 16, 20, 25, 30, 40, 50, 65, 85, 110, 150, 220, 350, 550, 900, 1500]

def get(url, tries=4):
    for a in range(tries):
        try:
            with urllib.request.urlopen(url, timeout=180) as r:
                return r.read().decode("utf-8")
        except Exception as e:  # noqa: BLE001 - retry any network error
            err = e
            time.sleep(3 + a * 5)
    raise err

wb = json.loads((HERE.parent / "wb_raw.json").read_text())
iso3to2 = {v["iso3"]: k for k, v in wb.items()}
cinfo = {c["id"]: c for c in json.loads((ROOT / "src/data/countries.json").read_text())}
ids = set(cinfo)

# PIP: one call per line, national rows only.
cdf, meta = {}, {}
for line in LINES:
    rows = csv.DictReader(io.StringIO(get(
        f"https://api.worldbank.org/pip/v1/pip?country=all&year={YEAR}&povline={line}&fill_gaps=true&format=csv")))
    for r in rows:
        iso2 = iso3to2.get(r["country_code"])
        if not iso2 or iso2 not in ids or r["reporting_level"] != "national" or not r["headcount"]:
            continue
        cdf.setdefault(iso2, []).append([line, round(float(r["headcount"]) * 100, 2)])
        meta[iso2] = {
            "med": round(float(r["median"]), 2) if r["median"] else None,
            "wt": "c" if r["welfare_type"] == "consumption" else "i",
            "y": int(r["reporting_year"]),
        }
    print("pip line", line, len(cdf), flush=True)

# WDI price level: latest year with both series.
def wdi(code):
    d = json.loads(get(f"https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&per_page=20000&date=2019:{YEAR}"))
    out = {}
    for row in d[1] or []:
        iso2 = iso3to2.get(row["countryiso3code"])
        if iso2 and row["value"] is not None:
            out.setdefault(iso2, {})[int(row["date"])] = float(row["value"])
    return out

prvt, fcrf = wdi("PA.NUS.PRVT.PP"), wdi("PA.NUS.FCRF")
# Guard against broken series (currency redenominations, official vs market
# exchange rates, zeros): a consumer price level older than two years, or more
# than 1.6 times off the economy-wide level (GDP per capita in USD / in PPP),
# is replaced by that economy-wide level and flagged as an estimate.
level, replaced = {}, []
for iso2 in ids:
    c = cinfo[iso2]
    gdp_level = c["gdp"] / c["ppp"] if c.get("gdp") and c.get("ppp") else None
    years = prvt.get(iso2, {})
    common = sorted(set(years) & set(fcrf.get(iso2, {})), reverse=True)
    pl = y = None
    if common and fcrf[iso2][common[0]] > 0:
        y = common[0]
        pl = years[y] / fcrf[iso2][y]
    ok = pl is not None and pl > 0 and y >= YEAR - 2 and (
        gdp_level is None or 1 / 1.6 <= pl / gdp_level <= 1.6)
    if ok:
        level[iso2] = {"pl": round(pl, 3), "plY": y}
    elif gdp_level:
        level[iso2] = {"pl": round(gdp_level, 3), "plY": YEAR, "plE": True}
        replaced.append(iso2)
print("price level", len(level), "from GDP", sorted(replaced))

# PIP lines are in 2021 PPP dollars; prices are in today's dollars.
cpi = json.loads(get("https://api.worldbank.org/v2/country/US/indicator/FP.CPI.TOTL?format=json&date=2021:2026"))[1]
cpi = {int(r["date"]): r["value"] for r in cpi if r["value"] is not None}
cpi_year = max(cpi)
us_deflator = round(cpi[cpi_year] / cpi[2021], 4)
print("US CPI", cpi_year, "vs 2021:", us_deflator)

# ITU: keep a small CSV extract in the repo so the build is reproducible.
extract = HERE / "itu_2025_extract.csv"
if len(sys.argv) > 1:
    import openpyxl
    ws = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)["economies_2008-2025"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = header.index(YEAR)
    keep = {"i271mb_high_5GB$", "i271mb_5GB$", "i154_FBB5$"}
    with extract.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["iso3", "code", "usd_2025"])
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[2] in keep and isinstance(row[col], (int, float)):
                w.writerow([row[0], row[2], row[col]])
itu = {}
for r in csv.DictReader(extract.open()):
    iso2 = iso3to2.get(r["iso3"])
    if iso2:
        itu.setdefault(iso2, {})[r["code"]] = float(r["usd_2025"])

out = {}
for iso2 in sorted(ids):
    rec = {}
    if iso2 in cdf and len(cdf[iso2]) == len(LINES):
        rec["cdf"] = [p for _, p in sorted(cdf[iso2])]
        rec.update(meta[iso2])
    rec.update(level.get(iso2, {}))
    t = itu.get(iso2, {})
    mob = t.get("i271mb_high_5GB$") or t.get("i271mb_5GB$")
    if mob:
        rec["mob"] = round(mob, 2)
    if t.get("i154_FBB5$"):
        rec["fbb"] = round(t["i154_FBB5$"], 2)
    if rec:
        out[iso2] = rec

result = {
    "sources": {
        "pip": {"title": "World Bank Poverty and Inequality Platform", "url": "https://pip.worldbank.org/", "license": "CC BY 4.0", "year": YEAR},
        "wdi": {"title": "World Bank WDI: PA.NUS.PRVT.PP / PA.NUS.FCRF", "url": "https://data.worldbank.org/indicator/PA.NUS.PRVT.PP", "license": "CC BY 4.0"},
        "itu": {"title": "ITU ICT Price Baskets, Dec 2025 release", "url": "https://www.itu.int/en/ITU-D/Statistics/Pages/ICTprices/default.aspx", "license": "CC BY-NC-SA 3.0 IGO", "year": YEAR},
    },
    "lines": LINES,
    # Divide a current-dollar PPP price by this to get 2021 PPP dollars (US CPI).
    "usDeflator": us_deflator,
    "usDeflatorYear": cpi_year,
    "countries": out,
}
(ROOT / "src/data/afford.json").write_text(json.dumps(result, separators=(",", ":")) + "\n")
n = lambda k: sum(1 for v in out.values() if k in v)
print(f"countries {len(out)}: cdf {n('cdf')}, price level {n('pl')}, mobile {n('mob')}, broadband {n('fbb')}")
print("RU", out.get("RU"))
